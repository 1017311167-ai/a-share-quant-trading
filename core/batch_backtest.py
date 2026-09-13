"""批量回测引擎 —— 多股票 × 多策略 并发回测

一次跑多只股票、多个策略的完整组合，用多进程并发加速，结果汇总成一张
DataFrame，支持按任意指标排序，并导出为 Excel（汇总表 + 每个组合明细）。

完全复用现有回测引擎（core/backtest_engine.py）和策略工厂
（strategies/factory.py）：每个组合内部走的就是单标的回测流程。

用法:
    from core.batch_backtest import run_batch, sort_results, export_excel

    out = run_batch(
        ["600519", "000001", "300750"],
        [{"name": "双均线", "params": {"fast": 5, "slow": 20}},
         {"name": "RSI超买超卖", "params": {}}],
        init_cash=1_000_000, commission=0.00025, slippage=0.001,
        max_workers=4)
    print(sort_results(out["results"], by="夏普比率"))
    export_excel(out, "data/批量回测结果.xlsx")

Windows 兼容说明:
    进程池用 ProcessPoolExecutor（Windows / macOS 默认都是 spawn 启动方式），
    所以：worker 函数必须是模块顶层函数（_run_one，可被 pickle），
    入口代码必须放在 if __name__ == "__main__": 保护下（run_test 已满足）。
"""

import datetime
import itertools
import logging
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed

# 保证直接运行本文件（python core/batch_backtest.py）时能找到 core / strategies / utils 包
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
from openpyxl.styles import Font

from core.backtest_engine import BacktestEngine
from core.risk_analysis import analyze_portfolio, trades_table
from strategies.factory import create_strategy
from utils.data_loader import load_market_data

logger = logging.getLogger("batch_backtest")

# 汇总表标准列（每个组合一行，含全部绩效指标）
RESULT_COLUMNS = [
    "股票代码", "策略名称", "参数",
    "累计收益率", "年化收益率", "年化波动率", "夏普比率", "卡玛比率",
    "最大回撤", "胜率", "盈亏比", "总交易次数", "总手续费", "期末总资产", "基准收益率",
]

# 可用于排序的指标列
SORTABLE = [c for c in RESULT_COLUMNS if c not in ("股票代码", "策略名称", "参数")]

# Excel 数字格式（汇总表 / 明细指标表共用）
_NUM_FMT = {
    "累计收益率": "0.00%", "年化收益率": "0.00%", "年化波动率": "0.00%",
    "最大回撤": "0.00%", "胜率": "0.00%", "基准收益率": "0.00%",
    "夏普比率": "0.00", "卡玛比率": "0.00", "盈亏比": "0.00",
    "总交易次数": "0", "总手续费": "#,##0.00", "期末总资产": "#,##0",
}


def _params_str(params: dict) -> str:
    """参数字典转展示字符串；空字典 -> "默认参数" """
    if not params:
        return "默认参数"
    return "、".join(f"{k}={v:g}" if isinstance(v, float) else f"{k}={v}"
                     for k, v in params.items())


def _run_one(task):
    """单个组合的回测（模块顶层函数，多进程 worker 直接调用）

    参数:
        task: (股票代码, 策略配置 dict, 行情 DataFrame, 引擎参数字典)

    返回:
        (是否成功, 结果行 dict 或 错误信息 str, 明细 dict 或 None)
    """
    code, cfg, df, engine_kwargs = task
    try:
        strategy = create_strategy(cfg["name"], **cfg.get("params", {}))
        entries, exits = strategy.generate_signals(df)
        if entries.sum() == 0:
            raise ValueError("策略没有产生任何买入信号")
        bt = BacktestEngine(df, entries, exits, code=code, **engine_kwargs).run()
        analyzer = analyze_portfolio(bt.portfolio, rf=engine_kwargs.get("rf", 0.0))
        m = bt.get_metrics()
        rm = analyzer.compute_metrics()
        row = {
            "股票代码": code,
            "策略名称": strategy.name,
            "参数": _params_str(cfg.get("params", {})),
            "累计收益率": m["累计收益率"],
            "年化收益率": m["年化收益率"],
            "年化波动率": rm["年化波动率"],
            "夏普比率": m["夏普比率"],
            "卡玛比率": rm["卡玛比率"],
            "最大回撤": m["最大回撤"],
            "胜率": m["胜率"],
            "盈亏比": rm["盈亏比"],
            "总交易次数": m["总交易次数"],
            "总手续费": m["总手续费"],
            "期末总资产": m["期末总资产"],
            "基准收益率": m["基准收益率"],
        }
        detail = {
            "metrics": row,
            "monthly": analyzer.monthly_returns(),
            "trades": trades_table(bt.portfolio),
        }
        return (True, row, detail)
    except Exception as e:
        return (False, f"{code} × {cfg['name']}（{_params_str(cfg.get('params', {}))}）: "
                       f"{type(e).__name__}: {e}", None)


def run_batch(codes, strategy_configs, init_cash=1_000_000, commission=0.00025,
              slippage=0.001, *, start=None, end=None, rf=0.0, freq="daily",
              t_plus_1=True, price_limit=True, stamp_tax=True, transfer_fee=True,
              commission_min=5.0, participation_rate=0.05,
              impact_coefficient=0.02, max_slippage=0.05,
              limit_queue_fill_ratio=0.25, order_ttl_bars=5,
              execution_model="realistic", trade_unit=100,
              max_workers=None, progress_cb=None, data_map=None,
              record_experiment: bool = True,
              experiment_name: str | None = None,
              experiment_store=None) -> dict:
    """多股票 × 多策略批量回测

    参数:
        codes:             股票代码列表，如 ["600519", "000001"]
        strategy_configs:  策略配置列表，每项 {"name": 策略名, "params": 参数字典}
                           （params 可省略，用策略默认参数）
        init_cash:         初始资金（元）
        commission:        佣金费率（小数，万 2.5 = 0.00025）
        slippage:          滑点（小数）
        start/end:         回测区间（默认近 3 年）
        freq:              数据频率："daily"（日线，默认）/
                           "1min" / "5min" / "15min" / "30min" / "60min"（分钟线）
        rf:                无风险利率（年化小数）
        t_plus_1 / price_limit / stamp_tax / transfer_fee: A股规则开关，同回测引擎
        max_workers:       并发进程数（默认 = min(组合数, CPU 核数)）
        progress_cb:       进度回调 progress_cb(已完成数, 总数)，供界面进度条使用
        data_map:          可选的已加载行情 {股票代码: DataFrame}，用于测试和重放
        record_experiment: 是否自动保存实验记录，默认 True
        experiment_name:   实验名称
        experiment_store:  自定义 ExperimentStore

    返回:
        {"results": 汇总 DataFrame（按 股票代码/策略名称/参数 排序）,
         "details": {组合key: {"metrics": 指标行, "monthly": 月度收益, "trades": 交易明细}},
         "failed": 失败组合及原因列表（异常标的自动跳过，不中断整体任务）,
         "elapsed": 总用时（秒）,
         "experiment_id": 自动记录的实验 ID,
         "reproducibility_key": 实验输入指纹,
         "result_hash": 实验结果哈希}
    """
    t0 = time.perf_counter()
    if not logger.handlers:
        logging.basicConfig(level=logging.INFO,
                            format="%(asctime)s [%(levelname)s] %(message)s",
                            datefmt="%H:%M:%S")
    if not codes:
        raise ValueError("股票代码列表不能为空")
    if not strategy_configs:
        raise ValueError("策略配置列表不能为空")

    if start is None or end is None:
        end = datetime.date.today()
        start = end - datetime.timedelta(days=3 * 365)

    failed = []

    # 1. 校验策略配置（提前把无效策略挑出来，不浪费进程）
    valid_configs = []
    for cfg in strategy_configs:
        try:
            create_strategy(cfg["name"])
        except Exception as e:
            logger.warning(f"策略配置无效，已跳过：{cfg!r}（{e}）")
            failed.append(f"策略配置无效：{cfg!r}（{e}）")
        else:
            valid_configs.append(cfg)
    if not valid_configs:
        raise ValueError("没有有效的策略配置，无法开始批量回测")

    # 2. 主进程顺序加载行情数据（避免多进程同时下载互相干扰）
    unit = "个交易日" if freq == "daily" else "根K线"
    dfs = {}
    if data_map is not None:
        for code in codes:
            if code in data_map:
                dfs[code] = data_map[code]
                logger.info(f"使用注入行情：{code}（{len(dfs[code])} {unit}）")
            else:
                failed.append(f"{code}：注入行情中缺少该股票")
    else:
        for code in codes:
            try:
                dfs[code] = load_market_data(code, start, end, freq=freq)
                logger.info(f"行情数据就绪：{code}（{len(dfs[code])} {unit}）")
            except Exception as e:
                logger.warning(f"{code} 行情数据加载失败，该股票全部组合跳过（{type(e).__name__}: {e}）")
                failed.append(f"{code}：行情数据加载失败（{type(e).__name__}: {e}）")

    engine_kwargs = {
        "init_cash": init_cash, "commission": commission, "slippage": slippage,
        "rf": rf, "t_plus_1": t_plus_1, "price_limit": price_limit,
        "stamp_tax": stamp_tax, "transfer_fee": transfer_fee,
        "commission_min": commission_min,
        "participation_rate": participation_rate,
        "impact_coefficient": impact_coefficient,
        "max_slippage": max_slippage,
        "limit_queue_fill_ratio": limit_queue_fill_ratio,
        "order_ttl_bars": order_ttl_bars,
        "execution_model": execution_model,
        "trade_unit": trade_unit,
    }
    tasks = [(code, cfg, dfs[code], engine_kwargs)
             for code, cfg in itertools.product(dfs.keys(), valid_configs)]

    total = len(tasks)
    if total == 0:
        logger.warning("所有股票都加载失败，没有可回测的组合")
        return {"results": pd.DataFrame(columns=RESULT_COLUMNS),
                "details": {}, "failed": failed, "elapsed": time.perf_counter() - t0}

    n_workers = max_workers or min(total, os.cpu_count() or 1)
    logger.info(f"开始批量回测：{len(dfs)} 只股票 × {len(valid_configs)} 个策略 "
                f"= {total} 个组合，并发进程数 {n_workers}")

    # 3. 多进程并发回测（每个组合一个任务，异常在 worker 内捕获，不中断整体）
    rows, details = [], {}
    done = 0
    with ProcessPoolExecutor(max_workers=n_workers) as pool:
        futures = [pool.submit(_run_one, task) for task in tasks]
        for fut in as_completed(futures):
            ok, row_or_err, detail = fut.result()
            done += 1
            if progress_cb:
                try:
                    progress_cb(done, total)
                except Exception:
                    pass  # 进度回调出错不影响主流程
            if ok:
                rows.append(row_or_err)
                key = f"{row_or_err['股票代码']}_{row_or_err['策略名称']}"
                base, i = key, 2
                while base in details:  # 同股同策略不同参数时 key 会撞名，加序号
                    base = f"{key}_{i}"
                    i += 1
                details[base] = detail
                logger.info(f"[{done}/{total}] ✓ {row_or_err['股票代码']} × "
                            f"{row_or_err['策略名称']}（{row_or_err['参数']}）："
                            f"累计 {row_or_err['累计收益率']:.2%}，"
                            f"夏普 {row_or_err['夏普比率']:.2f}")
            else:
                failed.append(row_or_err)
                logger.warning(f"[{done}/{total}] ✗ {row_or_err}")

    results = (pd.DataFrame(rows, columns=RESULT_COLUMNS)
               .sort_values(["股票代码", "策略名称", "参数"], kind="stable")
               .reset_index(drop=True))
    elapsed = time.perf_counter() - t0
    logger.info(f"批量回测完成：成功 {len(rows)} 个组合，失败 {len(failed)} 项，"
                f"用时 {elapsed:.1f} 秒")
    out = {"results": results, "details": details, "failed": failed, "elapsed": elapsed}
    if record_experiment:
        from research.experiments import CostAssumptions, record_batch_experiment

        record = record_batch_experiment(
            dfs,
            codes=list(codes),
            strategy_configs=valid_configs,
            output=out,
            costs=CostAssumptions(
                init_cash=init_cash,
                commission=commission,
                commission_min=commission_min,
                slippage=slippage,
                impact_coefficient=impact_coefficient,
                max_slippage=max_slippage,
                participation_rate=participation_rate,
                lot_size=trade_unit,
                stamp_tax=stamp_tax,
                transfer_fee=transfer_fee,
                t_plus_1=t_plus_1,
                price_limit=price_limit,
                limit_queue_fill_ratio=limit_queue_fill_ratio,
                order_ttl_bars=order_ttl_bars,
                execution_model=execution_model,
                rf=rf,
            ),
            start=start,
            end=end,
            frequency=freq,
            max_workers=max_workers,
            name=experiment_name,
            store=experiment_store,
        )
        out["experiment_id"] = record.experiment_id
        out["reproducibility_key"] = record.reproducibility_key
        out["result_hash"] = record.result_hash
    return out


def sort_results(results, by: str = "夏普比率", ascending=None) -> pd.DataFrame:
    """结果按指标排序（返回新表，不改原表）

    参数:
        results:   run_batch 返回的汇总 DataFrame
        by:        指标列名（夏普比率 / 年化收益率 / 最大回撤 等，见 SORTABLE）
        ascending: True 升序 / False 降序；None 时自动判断
                   （最大回撤越小越好 → 升序，其余指标越大越好 → 降序）

    返回:
        排序后的 DataFrame
    """
    if by not in results.columns:
        raise ValueError(f"未知排序指标：{by!r}，可选：{SORTABLE}")
    if ascending is None:
        ascending = by == "最大回撤"  # 回撤越小越好
    return results.sort_values(by, ascending=ascending, kind="stable").reset_index(drop=True)


# ---------------------------------------------------------------------------
# Excel 导出
# ---------------------------------------------------------------------------

def _safe_sheet_name(name: str, used: dict) -> str:
    """sheet 名去 Excel 非法字符、限长 31、重名自动加序号"""
    name = "".join("_" if ch in r"[]:*?/\\" else ch for ch in name).strip()
    name = (name or "组合")[:31]
    if name not in used:
        used[name] = 1
        return name
    used[name] += 1
    suffix = f" ({used[name]})"
    return name[:31 - len(suffix)] + suffix


def _write_block(writer, sheet, df, startrow, label=None):
    """在明细 sheet 里按块写入：可选加粗标题行 + DataFrame，返回下一块起始行"""
    if label:
        cell = writer.sheets[sheet].cell(row=startrow + 1, column=1, value=label)
        cell.font = Font(bold=True, size=12)
        startrow += 1
    df.to_excel(writer, sheet_name=sheet, index=False, startrow=startrow)
    return startrow + len(df) + 2  # 留一行空行


def _autofit(ws, max_width: int = 45):
    """按内容宽度自适应列宽（简单估算：最长字符数 × 1.8 + 2，中文按双宽）"""
    for col in ws.columns:
        width = max((len(str(c.value)) * (1.8 if any(ord(ch) > 127 for ch in str(c.value)) else 1.0)
                     for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, max_width)


def export_excel(out: dict, path: str) -> str:
    """把批量回测结果导出为 Excel

    结构:
        第 1 张 sheet「汇总」：所有组合的绩效指标总表（带数字格式）
        其余 sheet：每个组合一张明细（绩效指标 + 月度收益 + 交易明细）

    参数:
        out:  run_batch 的返回结果
        path: 导出路径（如 "data/批量回测结果.xlsx"）

    返回:
        path（可继续使用）
    """
    results = out.get("results")
    details = out.get("details", {})
    if results is None or len(results) == 0:
        raise ValueError("没有可导出的回测结果（results 为空）")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # 汇总表
        results.to_excel(writer, sheet_name="汇总", index=False)
        ws = writer.sheets["汇总"]
        header = {c: i + 1 for i, c in enumerate(results.columns)}
        for name, fmt in _NUM_FMT.items():
            if name in header:
                for row in range(2, len(results) + 2):
                    ws.cell(row=row, column=header[name]).number_format = fmt
        _autofit(ws)

        # 每个组合一张明细
        used = {}
        for key, detail in details.items():
            sheet = _safe_sheet_name(key, used)
            # 先建好并注册 sheet，才能在写数据前写标题行（pandas 按注册名复用）
            ws2 = writer.book.create_sheet(sheet)
            writer.sheets[sheet] = ws2
            metrics = detail.get("metrics", {})
            metrics_df = pd.DataFrame({"指标": list(metrics.keys()),
                                       "数值": list(metrics.values())})
            row = _write_block(writer, sheet, metrics_df, 0, label="一、绩效指标")
            for r in range(3, len(metrics_df) + 3):  # 数据从第 3 行开始（标题+表头占 2 行）
                name = ws2.cell(row=r, column=1).value
                if name in _NUM_FMT:
                    ws2.cell(row=r, column=2).number_format = _NUM_FMT[name]

            monthly = detail.get("monthly")
            if monthly is not None and len(monthly):
                mon = monthly.reset_index()
                mon.columns = ["月份", "月度收益率"]
                mon["月份"] = mon["月份"].dt.strftime("%Y-%m")
                block_start = row
                row = _write_block(writer, sheet, mon, block_start, label="二、月度收益")
                for r in range(block_start + 3, block_start + 3 + len(mon)):
                    ws2.cell(row=r, column=2).number_format = "0.00%"

            trades = detail.get("trades")
            if trades is not None and len(trades):
                _write_block(writer, sheet, trades, row, label="三、交易明细")
            _autofit(ws2)

        # openpyxl 工作簿默认带一张空的 "Sheet"，若存在则删除
        if "Sheet" in writer.book.sheetnames:
            del writer.book["Sheet"]
    return path


def run_test():
    """测试：3 只股票 × 2 策略批量回测 + 排序 + Excel 导出

    运行方式（在项目根目录下）：
        python core/batch_backtest.py
    """
    print("===== 批量回测测试开始 =====")
    codes = ["600519", "000001", "300750"]
    configs = [
        {"name": "双均线", "params": {}},
        {"name": "RSI超买超卖", "params": {}},
    ]
    out = run_batch(codes, configs, init_cash=500_000, commission=0.0003,
                    slippage=0.001, max_workers=3)
    results, failed = out["results"], out["failed"]

    # 1. 组合数 = 股票 × 策略（失败的进 failed 列表，不中断整体任务）
    assert len(results) + len(failed) == 6, \
        f"成功 {len(results)} + 失败 {len(failed)} 应等于 6"
    assert list(results.columns) == RESULT_COLUMNS, "汇总表列不完整"
    assert set(results["股票代码"]) <= set(codes)
    print(f"  ✓ 批量回测：成功 {len(results)} 个组合，失败 {len(failed)} 项，"
          f"并发 3 进程，用时 {out['elapsed']:.1f} 秒")
    for f in failed:
        print(f"    跳过：{f}")

    # 2. 600519 的两个策略都应成功（本地有缓存，不依赖网络）
    for name in ("双均线策略", "RSI超买超卖策略"):
        assert name in set(results["策略名称"]), f"600519 的 {name} 应在结果中"
    print("  ✓ 600519 × 双均线/RSI 均在结果中")

    # 3. 排序：夏普降序 / 最大回撤升序 / 未知指标报错
    by_sharpe = sort_results(results, by="夏普比率")
    assert (by_sharpe["夏普比率"].diff().dropna() <= 1e-9).all(), "夏普应降序排列"
    by_dd = sort_results(results, by="最大回撤")
    assert (by_dd["最大回撤"].diff().dropna() >= -1e-9).all(), "最大回撤应升序排列"
    try:
        sort_results(results, by="不存在的指标")
    except ValueError:
        pass
    else:
        raise AssertionError("未知排序指标应报错")
    print("  ✓ 排序：夏普降序 / 最大回撤升序 / 未知指标报错")

    # 4. Excel 导出 + 回读验证
    import tempfile

    from openpyxl import load_workbook
    path = os.path.join(tempfile.gettempdir(), "batch_backtest_test.xlsx")
    export_excel(out, path)
    wb = load_workbook(path)
    sheet_names = wb.sheetnames
    assert sheet_names[0] == "汇总"
    assert len(sheet_names) == 1 + len(out["details"]), "应有 1 张汇总 + 每个组合 1 张明细"
    assert wb[sheet_names[1]]["A1"].value == "一、绩效指标", "明细 sheet 应有指标块标题"
    assert len(pd.read_excel(path, sheet_name="汇总")) == len(results), "汇总表行数应一致"
    print(f"  ✓ Excel 导出：{len(sheet_names)} 个 sheet（1 汇总 + {len(out['details'])} 明细），"
          f"回读校验通过（{path}）")

    # 5. 打印排序后的汇总表（摘要列）
    display = ["股票代码", "策略名称", "参数", "累计收益率", "年化收益率", "夏普比率",
               "最大回撤", "胜率", "总交易次数", "总手续费"]
    print()
    print("---- 汇总表（按夏普比率降序） ----")
    print(by_sharpe[display].to_string(index=False, float_format=lambda x: f"{x:.4f}"))
    print()
    print("---- 汇总表（按最大回撤升序） ----")
    print(by_dd[display].to_string(index=False, float_format=lambda x: f"{x:.4f}"))
    print()
    print("===== 全部测试通过 =====")


if __name__ == "__main__":
    run_test()
